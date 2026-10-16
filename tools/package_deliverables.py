from pathlib import Path
import json,hashlib,shutil,zipfile
from openpyxl import load_workbook
import fitz
ROOT=Path('/workspace/assignment1');OUT=ROOT/'deliverables';R=Path('/workspace/assignment1-render')
shutil.copy2(R/'assignment1_report.pdf',OUT/'assignment1_report.pdf')
sources=json.loads((OUT/'sources.json').read_text())
lines=['# Official sources and claim mapping','','Audit date: 16 October 2026. Harvard author–date format. All commercial, cost, risk, workforce and behavioral parameters are illustrative assumptions or explicit derivations. Sources verify only the claims identified below.','']
evidence=[]
keys={'G1':'google2006.html','G2':'pricing.html','G3':'marketplace.html','G4':'gmail_page.html','G5':'directory_page.html','G6':'drive_page.html','M2':'ms_mailbox.txt','M3':'ms_migration.txt','M4':'online_description.html','M5':'architecture.html','M6':'exchangeSE.html'}
for s in sources:
 lines.extend([f'## {s["id"]} — {s["author"]} ({s["date"]})','',s['title']+'.','',f'Official source: {s["url"]}','',f'Supported use: {s["supports"]}','',f'Access audit: {s["status"]}',''])
 if s['mirror']:lines.extend([f'Official documentation/sample source: {s["mirror"]}',''])
 if s['id'] in keys:
  p=Path('/workspace/assignment1-research')/keys[s['id']]
  if p.exists():evidence.append({'source_id':s['id'],'retrieved_file':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'audit_date':'2026-10-16'})
lines.extend(['## Interpretation limits','','The Google page displays standard annual-commitment USD 7 and USD 14 levels alongside promotions and differing features. The five-year model assumes flat prices and excludes promotional discounts. No five-year guaranteed quote was obtained.','','The Microsoft commercial licensing overview returned HTTP 403. No exact licensing entitlement or price is claimed. The USD 12,000 annual allowance requires current SE, Windows/CAL and subscription/Software Assurance due diligence before procurement.','','Quickstarts were inspected as documentation only. No live tenant, credentials, customer migration or API endpoint was tested. No observed vendor CAC, churn, market-share or financial dataset was used.'])
(OUT/'official_sources.md').write_text('\n'.join(lines),encoding='utf-8')
(OUT/'source_access_evidence.json').write_text(json.dumps(evidence,indent=2))
pdf=fitz.open(OUT/'assignment1_report.pdf')
assert len(pdf)==9
for i,s in enumerate(['Executive Summary','Part A','Part B1','Part B2','Part B3','References','Appendix A','Appendix B','Appendix C']):assert s in pdf[i].get_text(),(i,s)
for page in pdf:assert abs(page.rect.width-595.3)<1 and abs(page.rect.height-841.9)<1
f=load_workbook(OUT/'assignment1_models.xlsx',data_only=False);cached=load_workbook(OUT/'assignment1_models.xlsx',data_only=True);r=load_workbook(R/'recalc/assignment1_models.xlsx',data_only=True)
count=0;diffs=[]
for sheet in f:
 for row in sheet:
  for c in row:
   if c.data_type!='f':continue
   count+=1;a=cached[sheet.title][c.coordinate].value;b=r[sheet.title][c.coordinate].value
   if isinstance(a,(int,float)) and isinstance(b,(int,float)):
    if abs(a-b)>1e-7*max(1,abs(a)):diffs.append((sheet.title,c.coordinate,a,b))
   elif a!=b:diffs.append((sheet.title,c.coordinate,a,b))
assert not diffs,diffs
checks=[r['Checks'].cell(row,2).value for row in range(4,17)]
assert checks==['PASS']*13,checks
assert not [(s.title,c.coordinate) for s in r for row in s for c in row if c.data_type=='e']
assert len(f['Charts']._charts)==2
w=load_workbook(R/'scenario-output/assignment1_models.xlsx',data_only=True)
scenario={'inputs':{'n':8,'google_price':14,'exchange_fte':.15},'ltv':w['CAC_LTV']['B8'].value,'exchange_tco':w['TCO']['B12'].value,'google_tco':w['TCO']['C12'].value,'google_savings':w['TCO']['B14'].value}
assert abs(scenario['ltv']-18678.042437064752)<.0001
assert scenario['exchange_tco']==221000 and scenario['google_tco']==248000 and scenario['google_savings']==-27000
v={'pdf_pages':9,'body_pages':5,'reference_pages':1,'appendix_pages':3,'page_size':'A4','body_columns':2,'body_font':'Arial 11 pt; tables/captions 10 pt','formula_cells':count,'independently_recalculated_baseline_checks_passed':13,'cached_vs_recalculated_differences':len(diffs),'formula_errors':0,'native_excel_charts':2,'input_propagation_check':scenario,'official_source_pages_verified':len(evidence),'pending_source':'M1: Microsoft commercial licensing overview returned HTTP 403; no licensing price or exact entitlement claim used','submission_status':'Not submitted; student VeriGuide receipt and official Signed Declaration not included'}
(OUT/'validation.json').write_text(json.dumps(v,indent=2))
readme='''# Assignment 1 交付说明

本套文件按你的案例制作，供本人审阅、修改后提交。

## 文件

- assignment1_report.docx：可编辑英文报告。
- assignment1_report.pdf：Word 导出的 PDF，共 9 页。第 1–5 页为 A4、双栏、单倍行距正文，第 6 页为参考文献，第 7–9 页为附录。正文 Arial 11 pt；表格与图注 10 pt。
- assignment1_models.xlsx：9 张工作表，包含所有输入、CAC/LTV、Churn、TCO、四种敏感性组合、两个原生 Excel 图表、来源与核对项。蓝色输入可修改。
- churn_chart / tco_chart：PNG、PDF、SVG 三种图表格式，均为初始基准情景的静态导出。
- official_sources.md / sources.csv：英文官方来源、支持的具体论点、核验状态。
- validation.json：分页、155 个公式独立重算、13 项基准核对及输入联动测试记录。
- assignment1_report.md：报告文字版。

## 主要口径

年度 Churn 12.06576083%；五年末留存 52.57630384%；五年 LTV USD 17,398.87；净分摊 CAC USD 7,000 / 4,000；LTV/CAC 2.49 / 4.35。

LTV 采用“当年贡献先发生、年末再流失”，现金在年末折现。若先计流失再计当年贡献，则为 USD 15,299.56，已提供对照。

交叉补贴用虚构 USD 30,000 的相邻业务资金抵扣解释。净分摊 CAC 下降不代表整个生态系统的真实资源成本减少；未声称 Google 实际采用该分摊。

五年 TCO：Exchange USD 281,000；Google USD 164,000；Google 少支出 USD 117,000（41.6%）。当 Google USD 14/月、Exchange 0.15 FTE 时，Exchange 便宜 USD 27,000。

所有商业、行为、风险与成本输入均标明 illustrative assumptions，未使用真实 CAC、流失率或客户数据。Google 网页的 USD 7/14 标准年承诺价已核对，但未来五年保持不变仍是假设。

## 来源与尚待事项

Google 定价、2006 年原公告、Marketplace 和 API 指南，以及 Microsoft Learn 的 Exchange SE、架构、Exchange Online 等来源已核验。Microsoft 商业授权网页仍返回 HTTP 403，已记录；USD 12,000/年只是许可预算。采购级结论仍需核对当前 SE、Windows/CAL、SA/订阅权益与报价。

## 提交前本人完成

1. 检查论证、文字、引用、课程对辅助工具的规则，并按要求补入姓名、学号等资料。
2. 完成学校规定的 VeriGuide，将真实 receipt 插入附录 C。
3. 使用学校官方 Signed Declaration，亲自完成并签署后插入附录。报告的说明页不能替代声明。
4. 重新导出最终 PDF，检查正文仍为五页，再提交 Blackboard。

附件列明截止时间：2026 年 10 月 16 日 17:00 HKT。当前文件未提交，未生成或代签任何证明。

## 修改与复现

Excel 编辑后需重新计算。Checks 的基准常数核对会提示合法输入变动；正文和外部图表不会随 Excel 自动更新，应同步修改。模型 horizon 固定为 5 年。

生成脚本位于仓库 tools/，依赖 python-docx、xlsxwriter、matplotlib、openpyxl、PyMuPDF 与 LibreOffice。脚本使用当前云环境的 /workspace/assignment1 路径；package_deliverables.py 还读取本次独立重算的临时验证文件。无需创建 Git worktree。原有 README 未修改。
'''
(OUT/'README_zh.md').write_text(readme,encoding='utf-8')
with zipfile.ZipFile(OUT/'assignment1_package.zip','w',compression=zipfile.ZIP_DEFLATED) as z:
 for p in sorted(OUT.iterdir()):
  if p.is_file() and p.suffix!='.zip':z.write(p,'assignment1/'+p.name)
 for p in sorted((ROOT/'tools').glob('*.py')):z.write(p,'assignment1/tools/'+p.name)
print(json.dumps(v,indent=2))
print('Bundle:',OUT/'assignment1_package.zip')
