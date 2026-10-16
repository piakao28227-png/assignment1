from pathlib import Path
import json, math, re
from docx import Document
from math_typesetting import omml
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path('/workspace/assignment1');OUT=ROOT/'deliverables'
M=json.loads((OUT/'model_results.json').read_text());SOURCES=json.loads((OUT/'sources.json').read_text())
D=Document();sec=D.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(1.55);sec.bottom_margin=Cm(1.5);sec.left_margin=Cm(1.55);sec.right_margin=Cm(1.55)
sec.header_distance=Cm(.6);sec.footer_distance=Cm(.65)
styles=D.styles
for name in ['Normal','Body Text','Caption']:
 styles[name].font.name='Arial';styles[name].font.size=Pt(11 if name!='Caption' else 10)
 styles[name].paragraph_format.line_spacing=1
 styles[name].paragraph_format.space_after=Pt(5)
for name in ['Heading 1','Heading 2','Heading 3']:
 styles[name].font.name='Arial';styles[name].font.size=Pt(11.5 if name=='Heading 1' else 11)
 styles[name].font.color.rgb=RGBColor.from_string('183B56');styles[name].font.bold=True
 styles[name].paragraph_format.space_before=Pt(7);styles[name].paragraph_format.space_after=Pt(4)
 styles[name].paragraph_format.keep_with_next=True
styles['Title'].font.name='Arial';styles['Title'].font.size=Pt(12);styles['Title'].font.bold=True
styles['Title'].font.color.rgb=RGBColor.from_string('183B56');styles['Title'].paragraph_format.space_after=Pt(7)
styles['Normal'].paragraph_format.widow_control=True
styles['Caption'].font.color.rgb=RGBColor.from_string('5B6B78')

def cols(section,n):
 el=section._sectPr.find(qn('w:cols'))
 if el is None:el=OxmlElement('w:cols');section._sectPr.append(el)
 el.set(qn('w:num'),str(n));el.set(qn('w:space'),'360')
cols(sec,2)
header=sec.header.paragraphs[0];header.text='ASSIGNMENT 1  |  STRATEGY ANALYSIS  |  2026–2030'
header.runs[0].font.name='Arial';header.runs[0].font.size=Pt(9);header.runs[0].font.color.rgb=RGBColor.from_string('5B6B78')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Assignment 1  |  ');r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
mathpr=D.settings.element.find(qn('m:mathPr'))
if mathpr is None:mathpr=OxmlElement('m:mathPr');D.settings.element.append(mathpr)
mathfont=mathpr.find(qn('m:mathFont'))
if mathfont is None:mathfont=OxmlElement('m:mathFont');mathpr.append(mathfont)
mathfont.set(qn('m:val'),'Latin Modern Math')
md=[]
pending_break=False
def use_break(pp):
 global pending_break
 if pending_break:pp.paragraph_format.page_break_before=True;pending_break=False
def h(text,level=1):
 pp=D.add_paragraph(text,f'Heading {level}');use_break(pp);md.append('#'*(level+1)+' '+text+'\n')
def cite_update(text):
 replacements={
  '(Google, n.d.-c; n.d.-d; n.d.-e; Microsoft, n.d.-c)':'(Google, 2026a, 2026b, 2026c; Microsoft, 2026b)',
  '(Google, n.d.-c; n.d.-d; n.d.-e)':'(Google, 2026a, 2026b, 2026c)',
  '(Microsoft, n.d.-b; n.d.-c)':'(Microsoft, 2026a, 2026b)',
  '(Google, n.d.-a)':'(Google, n.d.)',
  '(Google, n.d.-b)':'(Google, 2025)',
  '(Microsoft, n.d.-d)':'(Microsoft, 2025a)',
  '(Microsoft, n.d.-e)':'(Microsoft, 2025b)',
  '(Microsoft, n.d.-f)':'(Microsoft, 2025c)'}
 for old,new in replacements.items():text=text.replace(old,new)
 return text

def inline(pp,text):
 for piece in re.split(r'(\$[^$]+\$)',text):
  if piece.startswith('$') and piece.endswith('$'):
   pp._p.append(omml(piece[1:-1]))
  elif piece:pp.add_run(piece)

def p(text,style=None):
 text=cite_update(text)
 pp=D.add_paragraph(style=style);use_break(pp);inline(pp,text)
 pp.paragraph_format.line_spacing=1;md.append(text+'\n');return pp

def equation(latex):
 pp=D.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pp.paragraph_format.space_before=Pt(3);pp.paragraph_format.space_after=Pt(5)
 pp._p.append(omml(latex,display=True))
 md.append('$$\n'+latex+'\n$$\n')
 return pp

def table(heads,rows,widths=None):
 t=D.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.LEFT;t.autofit=False
 if widths:
  for column,w in zip(t.columns,widths):column.width=Cm(w)
  tw=t._tbl.tblPr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(int(Cm(sum(widths)).twips)))
 if widths:
  for cell,w in zip(t.rows[0].cells,widths):cell.width=Cm(w)
 for cell,s in zip(t.rows[0].cells,heads):cell.text=s
 for row in rows:
  for cell,s in zip(t.add_row().cells,row):cell.text=str(s)
 for i,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();n=OxmlElement('w:cantSplit');trPr.append(n)
  for j,cell in enumerate(row.cells):
   if widths:cell.width=Cm(widths[j])
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
   tcPr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side,val in [('top',55),('bottom',55),('left',65),('right',65)]:
    node=OxmlElement('w:'+side);node.set(qn('w:w'),str(val));node.set(qn('w:type'),'dxa');mar.append(node)
   tcPr.append(mar)
   if i==0 or i%2==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'183B56' if i==0 else 'EEF5F8');tcPr.append(shade)
   for pp in cell.paragraphs:
    pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.line_spacing=1
    for run in pp.runs:
     run.font.name='Arial';run.font.size=Pt(10)
     if i==0:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
 md.append('| '+' | '.join(heads)+' |\n| '+' | '.join(['---']*len(heads))+' |\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in rows)+'\n')
 D.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

def colbreak():D.add_paragraph().add_run().add_break(WD_BREAK.COLUMN);md.append('\n<!-- column break -->\n')
def pagebreak():
 global pending_break
 pending_break=True;md.append('\n---\n')
def figure(name,caption):
 pp=D.add_paragraph();pp.paragraph_format.space_after=Pt(3);pp.add_run().add_picture(str(OUT/name),width=Cm(8.5))
 p(caption,'Caption');md.append(f'![{caption}]({name})\n')

# Page 1
p('From On-Premises Email to Cloud Platforms: A Strategic and Quantitative Analysis of Microsoft Exchange and Google Workspace','Title')
h('Executive Summary')
p('For a hypothetical 200-employee SME, Google Workspace has the lower five-year cash cost in the baseline: USD 164,000 versus USD 281,000 for Exchange Server Subscription Edition (SE). The modeled saving is USD 117,000, or 41.6%, over 2026–2030. This is a conditional planning result; it is not an observed saving, a vendor quotation or a recommendation independent of functional requirements.')
p('A standardized supplier-side account model produces five-year risk-adjusted lifetime value (LTV) of USD 17,399 for either option. Against illustrative customer acquisition costs (CAC) of USD 7,000 for Exchange and USD 4,000 for Google, the ratios are 2.49 and 4.35. The difference comes entirely from assumed acquisition-cost allocation, not superior measured Google retention or margins.')
p('The modeled annual churn rate declines from 18.00% at zero integrations to 8.09% at eight. At the baseline of four integrations, churn is 12.07% and five-year end retention is 52.58%. These are outputs of an assumed functional relationship, not customer-cohort estimates.')
h('Decision in brief',2)
p('Shortlist Google Workspace for this cost-sensitive case, while obtaining a like-for-like Exchange Online quote and testing migration fidelity, security and governance. Cost leadership is fragile: at USD 14/user/month and only 0.15 Exchange FTE, Exchange SE becomes USD 27,000 cheaper.')
colbreak()
h('Case Background')
p('The incumbent model places messaging infrastructure under the customer’s operational control. Exchange Mailbox servers store and process mailbox data, with administration through the Exchange admin center and Exchange Management Shell (Microsoft, n.d.-e). For the forward-looking 2026–2030 comparison, the on-premises option is Exchange SE, rather than an assumed deployment of an older server edition (Microsoft, n.d.-f).')
p('Google entered hosted organizational applications in 2006. Its 28 August announcement introduced Google Apps for Your Domain, combining Gmail, Google Talk, Google Calendar and Google Page Creator without customer-installed hardware or software (Google, 2006). This earlier offering provides the historical entry point; the modern comparison uses Google Workspace, not the 2006 bundle.')
p('The present Workspace pricing page describes business email and a wider productivity suite, with differentiated storage, meeting and management capabilities (Google, n.d.-a). Exchange SE primarily anchors this case’s messaging workload. The analysis therefore compares ways to meet a defined SME communications need, rather than asserting that the products contain identical features.')
h('Scope and evidence',2)
p('Part A evaluates platform mechanisms and industry structure. Part B1 takes the supplier’s perspective per enterprise account; B2 connects an integration scenario to churn and LTV; B3 takes the SME buyer’s perspective across 200 seats. Customer contribution is not the buyer’s cost saving.')
p('All commercial and behavioral model inputs are illustrative assumptions. Official documentation supports product and API descriptions, rather than estimates of vendor CAC, churn or market share. The TCO figures are planning budgets; Section B3 and Appendix C identify the entitlement, quotation and operational checks required before procurement.')
pagebreak()
# Page 2
h('Part A | Strategic Analysis')
h('Which platform effects matter?',2)
p('A user ecosystem can create direct coordination value: adoption by colleagues and business partners may make shared calendars, permissions and collaboration easier. However, email interoperability limits the need for external correspondents to use the same supplier. A larger email subscriber count alone is insufficient evidence of a strong proprietary network effect.')
p('Product complementarity is distinct. Email, storage, documents and meetings may be more convenient together than separately; this is an economy of scope or bundle benefit even if the number of users is unchanged. Workspace’s documented bundle illustrates the complement set, but does not measure the incremental value of each component (Google, n.d.-a).')
p('Third-party applications can generate an indirect network effect: more relevant customers attract developers; more useful integrations can then attract customers. Google documents third-party application discovery and installation through Workspace Marketplace (Google, n.d.-b). This establishes a platform channel, not proof of a self-reinforcing growth rate. Microsoft also exposes management and migration capabilities across SE and Exchange Online (Microsoft, n.d.-b; n.d.-c).')
p('Spreading datacenter, engineering or support expenditure over more customers is a supply-side scale economy. It can lower average cost without raising any individual customer’s utility through other users. It should not be counted again as a demand-side network effect.')
h('Bidirectional switching costs',2)
p('Leaving Exchange can require mailbox and calendar conversion, permission mapping, identity redesign, application rewrites and staff retraining. Leaving Workspace can similarly require document conversion, sharing-rule reconstruction, OAuth reauthorization and replacement of cloud workflows. Migration commands and API credentials provide concrete reasons for these tasks, but do not quantify their cost (Google, n.d.-c; n.d.-d; n.d.-e; Microsoft, n.d.-c).')
colbreak()
h('Five Forces comparison matrix',2)
table(['Force','Traditional on-premises model','Cloud-platform entry'],[
 ('Rivalry','Installed workflows, control and local service shape competition.','Suite bundling and rapid service updates broaden rivalry.'),
 ('New entrants','Deployment skills and customer infrastructure create entry friction.','Hosted delivery reduces deployment friction; trust and ecosystem depth remain barriers.'),
 ('Buyer power','Custom integrations and sunk setup costs weaken easy switching.','Visible subscriptions aid comparison; data and workflow dependence can weaken later bargaining.'),
 ('Supplier power','Customer depends on hardware, licensing and specialist support.','Some dependence shifts to the platform’s prices, APIs and service terms.'),
 ('Substitutes','Other mail servers or outsourced hosting can meet messaging needs.','Cloud suites and collaboration tools can substitute some workflows, while email remains useful.')
 ],[2.05,3.23,3.22])
p('These are directional analytical judgments, not empirical force scores or market-share estimates. Entry changes the location of bargaining power: initial deployment becomes easier, while a mature customer may face greater dependency on identity, data formats and complementary applications.')
h('The incumbent can also adapt',2)
p('Microsoft Exchange Online is itself a hosted messaging service, integrated with Microsoft Entra ID (Microsoft, n.d.-d). Its presence prevents a simple “Google cloud versus Microsoft servers” conclusion. An SME with existing Microsoft skills or workflows may prefer Microsoft’s cloud route and avoid part of the cross-vendor conversion burden.')
p('APIs have two effects. Proprietary workflows and permissions can increase rebuilding effort, but documented interfaces and open standards can support extraction and portability. The strategic recommendation is to evaluate workflow depth and exit feasibility together, rather than treating every integration as permanent lock-in.')
pagebreak()
# Page 3
h('Part B1 | CAC & Risk-Adjusted LTV')
h('An explicit cross-subsidy scenario',2)
p('Assume each supplier acquires 10 enterprise accounts. Let $A_j$ be acquisition spending allocated to supplier $j$ and $N_j$ the number of acquired accounts.')
equation('\\mathrm{CAC}_j=\\frac{A_j}{N_j}')
p('Exchange allocates USD 70,000, giving $\\mathrm{CAC}_{\\mathrm{E}}=\\frac{70{,}000}{10}=7{,}000$. For Google, assume a USD 70,000 gross campaign and a USD 30,000 ecosystem funding credit. Net allocated spending is USD 40,000, giving $\\mathrm{CAC}_{\\mathrm{G}}=\\frac{40{,}000}{10}=4{,}000$ per account.')
p('The illustrative credit shows how a diversified entrant could fund acquisition from adjacent offerings while assigning less expense to one product. It is a scenario assumption, rather than evidence of a Google accounting practice or advertising subsidy. Any reduction in gross acquisition expense through shared distribution would be a separate efficiency.')
p('A funding transfer does not remove real resource use: Google’s gross ecosystem acquisition cost remains USD 7,000 per account in this scenario. Therefore 4.35 is a ratio to net allocated CAC; on a common gross-cost basis both suppliers have a ratio of 2.49. Any claim of superior economic efficiency would need actual incremental cost evidence.')
h('Comparable contribution and risk',2)
p('Normalize annual contribution gross profit to USD 6,000 per active enterprise account for both suppliers. Apply a 5% expected operating-risk loss to contribution and a 10% annual discount rate. The loss haircut is distinct from churn and is applied once to each year’s contribution; it is not an extra customer-exit probability.')
p('Use the unrounded B2 baseline:')
equation('c(4)=0.18\\,\\mathrm{e}^{-0.4}\\approx0.1206576083.')
p('Equal churn and contribution isolate the assumed CAC difference. Acquisition expense is excluded from LTV because CAC appears separately in the ratio.')
colbreak()
h('Timing, formula and result',2)
p('A customer active at the start of a year earns that year’s contribution; churn occurs at year-end. Cash is received and discounted at year-end. Hence contribution eligibility in year $t$ is $(1-c)^{t-1}$. There is no terminal value or addition of new accounts.')
equation('\\mathrm{LTV}_5=\\sum_{t=1}^{5}\\frac{g(1-\\ell)(1-c)^{t-1}}{(1+r)^t}')
p('Here $g=6{,}000$ is annual contribution in USD, $\\ell=0.05$ is the operating-loss fraction and $r=0.10$ is the discount rate. With baseline $c=c(4)$, $\\mathrm{LTV}_5=17{,}398.87$ USD.')
table(['Calendar year','Start-year survival','PV contribution (USD)'],[(str(2025+t),f'{(1-M["baseline_churn"])**(t-1):.2%}',f'{6000*.95*(1-M["baseline_churn"])**(t-1)/1.1**t:,.2f}') for t in range(1,6)],[2.0,2.7,3.8])
table(['Measure','Exchange','Google'],[
 ('CAC, USD/account','7,000','4,000'),
 ('Five-year LTV, USD','17,398.87','17,398.87'),
 ('LTV / net CAC','2.49','4.35'),
 ('LTV / gross CAC','2.49','2.49')
 ],[3.5,2.5,2.5])
h('Interpretation and timing check',2)
p('The five-year finite horizon fits the case and avoids extrapolating indefinite retention. Ratios above one indicate positive modeled contribution coverage of acquisition cost before omitted fixed costs; they do not establish company profitability.')
p('If churn instead occurs before each year’s contribution, including year 1, multiply the baseline LTV by (1 − c): USD 15,299.56. The difference is a timing assumption, not a calculation error. The workbook reports both conventions and links the baseline directly to B2.')
pagebreak()
# Page 4
h('Part B2 | Switching Cost & Churn')
h('Integration-based behavioral model',2)
p('Define n as the count of active, material API/workflow integrations per enterprise account. The five values are alternative configurations, not observations across time. Hold n and annual churn constant within each configuration during 2026–2030. An integration must represent a meaningful dependency; merely enabling an API does not create the assumed effect.')
p('Annual churn follows the illustrative exponential function')
equation('c(n)=c_0\\mathrm{e}^{-\\beta n}=0.18\\,\\mathrm{e}^{-0.10n}.')
p('The baseline $c_0=0.18$ and coefficient $\\beta=0.10$ are assumed, not statistically fitted. Each additional integration multiplies churn by $\\mathrm{e}^{-0.10}\\approx0.9048$, a 9.52% proportional reduction under this model.')
p('For an account present at opening-2026, the probability of retention at end-2030 is')
equation('S_5(n)=\\left[1-c(n)\\right]^5.')
p('Annual exits are independent, with no re-entry and fixed $n$ within a scenario. End-horizon retention uses exponent 5; B1 contribution eligibility uses $t-1$ because contribution precedes year-end churn.')
table(['n','Annual churn','End-2030 retention'],[(str(s['n']),f'{s["churn"]:.2%}',f'{s["retention5"]:.2%}') for s in M['churn_scenarios']],[1.0,3.2,4.3])
h('Feedback to LTV',2)
p('Substituting each churn rate into B1 increases five-year LTV from USD 15,671 at n = 0 to USD 18,678 at n = 8, with all other inputs held fixed. At n = 4, the common LTV is USD 17,399. The workbook exposes every scenario and both CAC ratios, rather than hard-coding a displayed 12.07%.')
colbreak()
figure('churn_chart.png','Figure 1. Modeled annual churn by integration count; illustrative 2026–2030 configurations. Source: author’s scenario calculations; units shown on axes.')
h('Why APIs can lock in and open up',2)
p('Google’s official quickstarts show OAuth scopes and credentials for Gmail, Directory user administration and Drive metadata access (Google, n.d.-c; n.d.-d; n.d.-e). Recreating permissions, service identities and business logic on another platform can be costly. These concrete dependencies motivate the model’s proposed mechanism; API documentation does not establish the coefficient.')
p('The same interfaces can reduce manual extraction and reconciliation work. Portable formats and supported standards can make switching easier, while platform-specific permissions, histories and automation can remain difficult to reproduce. A functioning metadata-list sample is not proof of full-fidelity data export or a tested migration.')
h('What the model cannot establish',2)
p('A causal estimate would require account-level integration histories, customer size, service quality, prices and observed exits, with controls for customer selection. Loyal customers may integrate more rather than integrations causing loyalty. Satisfaction, contractual terms and new portability rules may also alter churn.')
p('Apply the function symmetrically to both suppliers in B1. It does not show that Google’s actual churn is below Microsoft’s, or that eight integrations are optimal. The SME should test exit procedures and migration fidelity alongside integration value.')
pagebreak()
# Page 5
h('Part B3 | Five-Year TCO')
h('Cost boundary and baseline',2)
p('Use nominal USD cash costs for 200 employees over 2026–2030. One-off payments occur at opening-2026, followed by five annual recurring payments. Users and prices remain constant, with no inflation, discounting, tax, hardware refresh or residual value. This cash-budget TCO is intentionally distinct from discounted supplier LTV.')
table(['Illustrative cost (USD)','Exchange','Google'],[
 ('Hardware/storage, once','32,000','—'),
 ('Deployment / migration, once','9,000','15,000'),
 ('License / subscription, yearly','12,000','16,800'),
 ('Administration, yearly','24,000','8,000'),
 ('Facilities etc., yearly','7,000','—'),
 ('Backup / security, yearly','5,000','5,000'),
 ('Five-year total','281,000','164,000')
 ],[4.15,2.18,2.17])
p('Applying $\\mathrm{TCO}=C_0+\\sum_{t=1}^{5}C_t$, the nominal totals in USD are')
equation('\\mathrm{TCO}_{\\mathrm{E}}=41{,}000+5(48{,}000)=281{,}000,')
equation('\\mathrm{TCO}_{\\mathrm{G}}=15{,}000+5(29{,}800)=164{,}000.')
p('Google saves USD 117,000; the proportional reduction is $\\frac{117{,}000}{281{,}000}\\times100\\%=41.6\\%$.')
p('USD 7 and USD 14 are illustrative flat subscription scenarios across five years. The retrieved Google page displays standard annual-commitment price points at these levels, alongside promotions and differentiated plans (Google, n.d.-a); it does not guarantee the modeled five-year rates. All hardware, licensing and labor amounts remain assumptions, not official quotations.')
p('Common endpoints and internet are excluded equally; USD 7,000 represents incremental on-premises overhead. Existing directory/DNS capacity is assumed reusable. The two-server allowance is not a validated resilient design (Microsoft, 2025b). Before procurement, confirm Exchange SE, Exchange Server/CAL and Windows Server licensing, SA or subscription eligibility, and current like-for-like quotations. The USD 12,000 annual license allowance is a planning assumption.')
colbreak()
figure('tco_chart.png','Figure 2. Five-year nominal TCO, 200 users, 2026–2030. Baseline: Google USD 7/user/month; Exchange 0.30 FTE. Source: author’s illustrative budget model.')
h('Sensitivity: when the ranking reverses',2)
p('The implied loaded annual FTE cost is $\\frac{24{,}000}{0.30}=80{,}000$ USD. At 0.15 FTE, Exchange administration falls to USD 12,000/year and TCO to USD 221,000. Raising Google’s price to USD 14 increases its TCO to USD 248,000.')
table(['Google price / Exchange FTE','Google savings (USD)'],[
 ('$7 / 0.30','117,000'),('$7 / 0.15','57,000'),('$14 / 0.30','33,000'),('$14 / 0.15','−27,000')
 ],[5.1,3.4])
p('Positive savings favor Google; the final case favors Exchange. Break-even Google prices are USD 16.75/user/month at 0.30 Exchange FTE and USD 11.75 at 0.15 FTE. Different Google plans may add capabilities, so this is a cost stress rather than a feature-equivalent comparison.')
h('Conclusion',2)
p('Google is the lower-cost baseline choice for this SME, conditional on successful migration and adequate functionality. Before procurement, test mailbox/calendar fidelity, identities, backup restoration and data-governance requirements, then compare current like-for-like Exchange SE, Google Workspace and Exchange Online quotes. The result supports a disciplined shortlist, not a universal claim that cloud or Google always wins.')

# References outside the five-page body, single-column.
sec2=D.add_section(WD_SECTION_START.NEW_PAGE);cols(sec2,1)
h('References')
for s in sorted(SOURCES,key=lambda s:(s['author'],s['date'],s['title'])):
 text=f'{s["author"]} ({s["date"]}) {s["title"]}. Available at: {s["url"]}.'
 text+=' (Accessed: 16 October 2026).'

 pp=p(text);pp.paragraph_format.space_after=Pt(7)
pagebreak()
h('Appendix A | Full Input Register')
p('All values below are illustrative assumptions, not measured vendor economics or procurement quotations. The workbook’s Inputs sheet is the editable source of model parameters. USD is used throughout.')
table(['Model / parameter','Value','Interpretation / boundary'],[
 ('Case / licensed users','200','Constant during 2026–2030.'),
 ('Horizon / first calendar year','5 years / 2026','No terminal value; five annual periods.'),
 ('B1 / acquired accounts','10 per supplier','Separate supplier-side acquisition scenario.'),
 ('B1 / Exchange acquisition spending','USD 70,000','Net and gross acquisition budget.'),
 ('B1 / Google gross spending / credit','USD 70,000 / 30,000','Hypothetical ecosystem-funded allocation credit.'),
 ('B1 / Google net spending','USD 40,000','Derived; gross resource spending remains USD 70,000.'),
 ('B1 / contribution per account','USD 6,000/year','Same for both suppliers; before risk haircut.'),
 ('B1 / expected risk loss / discount','5% / 10% per year','Separate operating-risk haircut and time value.'),
 ('B2 / c0 / beta / baseline n','0.18 / 0.10 / 4','Assumed churn function; not fitted to observations.'),
 ('B2 / integration scenarios','0, 2, 4, 6, 8','Material dependencies per enterprise account.'),
 ('B3 / servers and storage','USD 32,000 once','Two-server cost allowance, not an engineered design.'),
 ('B3 / Exchange deployment','USD 9,000 once','Opening-2026.'),
 ('B3 / licensing and SA allowance','USD 12,000/year','Exchange and underlying OS/CAL scope must be quoted.'),
 ('B3 / loaded FTE / baseline staffing','USD 80,000/year / 0.30','Derived annual administration cost USD 24,000.'),
 ('B3 / Exchange overhead','USD 7,000/year','Incremental facilities/network/maintenance.'),
 ('B3 / backup and added security','USD 5,000/year/option','Retained separately for each alternative.'),
 ('B3 / cloud migration and training','USD 15,000 once','Opening-2026.'),
 ('B3 / Google subscription','USD 7/user/month','Flat five-year rate; USD 16,800/year at 200 users.'),
 ('B3 / Google administration','USD 8,000/year','Incremental management allowance.'),
 ('Stress / Google price / Exchange staff','USD 14 / 0.15 FTE','One or both changes; all else held constant.')
 ],[6.1,4.1,7.7])
h('Scope of the comparison',2)
p('The baseline is a planning comparison of messaging-centered communications, not an assertion of identical productivity suites, security policies or service levels. No cloud productivity benefit, avoided outage cost or server resale value is monetized. No hardware refresh, extra directory-server purchase, Exchange Online TCO or network upgrade is estimated. Requirements that force a higher plan or additional infrastructure must be added before a procurement decision.')
pagebreak()
h('Appendix B | Calculation Detail & Audit')
h('B1–B2 expansion and time convention',2)
p('At the baseline $n=4$, calculations retain unrounded values:')
equation('c(4)=0.18\\,\\mathrm{e}^{-0.4}=0.1206576083\\ldots')
equation('g(1-\\ell)=6{,}000(1-0.05)=5{,}700\\quad\\text{USD/year}.')
p('For an account active at opening-2026, the present value of year-$t$ contribution and the resulting five-year total are')
equation('\\mathrm{PV}_t=\\frac{5{,}700(1-c)^{t-1}}{1.10^t},\\qquad \\mathrm{LTV}_5=\\sum_{t=1}^{5}\\mathrm{PV}_t=17{,}398.8669\\ldots')
equation('S_5(4)=(1-c)^5=0.5257630384\\ldots')
table(['n','Annual churn','End-2030 retention','LTV, USD'],[(str(s['n']),f'{s["churn"]:.6%}',f'{s["retention5"]:.6%}',f'{s["ltv"]:,.2f}') for s in M['churn_scenarios']],[2,4.3,5.5,6.1])
h('B3 sensitivity and break-even price',2)
table(['Google USD/user/month','Exchange FTE','Exchange TCO, USD','Google TCO, USD','Google savings, USD'],[(str(s['price']),f'{s["fte"]:.2f}',f'{s["exchange"]:,.0f}',f'{s["google"]:,.0f}',f'{s["saving"]:,.0f}') for s in M['sensitivity']],[3.6,2.6,4.0,3.8,3.9])
p('Let $f$ denote Exchange staffing in FTE and $p$ the Google subscription price in USD/user/month. Holding other inputs fixed gives')
equation('E(f)=41{,}000+5\\left(12{,}000+80{,}000f+7{,}000+5{,}000\\right),')
equation('G(p)=15{,}000+5\\left(2{,}400p+8{,}000+5{,}000\\right).')
p('Set $E(f)=G(p)$ and solve for the break-even subscription price:')
equation('p^{*}(f)=\\frac{E(f)-15{,}000-5(13{,}000)}{12{,}000}.')
p('Thus $p^{*}(0.30)=16.75$ and $p^{*}(0.15)=11.75$, in USD/user/month.')
h('Workbook structure and reconciliation',2)
p('Inputs contains the scenario assumptions; CAC_LTV and Churn link retention to discounted contribution; TCO reconciles category totals with annual cash flows; Sensitivity calculates four combinations and break-even prices. Charts contains linked Excel figures. Sources maps claims to documentation. Checks verifies arithmetic and model links, rather than the empirical validity of the assumptions.')
h('Boundaries of risk adjustment',2)
p('The 5% haircut applies to retained-account contribution and excludes exit risk already captured by churn. Platform-wide correlated losses are not modeled. If year-1 contribution requires surviving year-1 churn, five-year LTV falls to USD 15,299.56. Both conventions discount five end-year payments.')
pagebreak()
h('Appendix C | Procurement Scope & Submission Evidence')
h('Procurement qualification',2)
p('This case provides a five-year planning comparison, not a binding procurement quotation. All hardware, licensing, staffing, backup and migration costs are illustrative budgets. The standard Google price points provide market context; constant prices, staffing and seat counts through 2030 remain explicit modeling assumptions.')
p('Before procurement, verify that the USD 12,000 annual allowance covers the required Exchange SE entitlement, Exchange Server and Client Access Licenses (CALs), underlying Windows Server licensing, and applicable Software Assurance (SA) or subscription eligibility. Establish the organization’s existing entitlements and obtain dated, like-for-like supplier or reseller quotations. Add any uncovered licensing and infrastructure costs to the model.')
p('Technical due diligence should test mailbox and calendar migration fidelity, identity and permission mapping, backup restoration, continuity and data-governance requirements. Validate server sizing and resilience before accepting the two-server allowance. Compare a requirements-matched Exchange Online option with both modeled alternatives. No live tenant test or customer migration is claimed.')
h('Submission evidence',2)
p('The final submission must include the institution’s authentic VeriGuide receipt and completed Signed Declaration. The submitting student should review the analysis, calculations and citations and supply any required name, student ID and course information.')
p('The supplied assignment brief specifies Blackboard submission by 16 October 2026, 17:00 HKT. Insert the required receipt and declaration in this appendix before exporting the final PDF.')
table(['Required item','Status / student action'],[
 ('VeriGuide receipt','NOT INCLUDED. Complete the institutional check and insert the authentic receipt here.'),
 ('Signed Declaration','NOT INCLUDED. Complete and sign the official declaration; do not treat this page as the declaration.'),
 ],[5.3,12.6])
p('References and appendices are excluded from the five-page body limit under the supplied assignment brief.')

D.core_properties.title='From On-Premises Email to Cloud Platforms'
D.core_properties.subject='Assignment 1: Strategy Analysis — illustrative Exchange SE and Google Workspace case'
D.core_properties.author='Student review draft'
D.core_properties.keywords='Exchange SE, Google Workspace, CAC, LTV, churn, TCO, illustrative assumptions'
D.save(OUT/'assignment1_report.docx')
(OUT/'assignment1_report.md').write_text('\n'.join(md),encoding='utf-8')
print('Wrote report DOCX and Markdown:',OUT)
